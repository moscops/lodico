import unittest
from datetime import date, datetime, time, timedelta
from io import BytesIO

from openpyxl import load_workbook

from bot.config import ROOT, load_config
from bot.db import Database
from bot.planner import available_between, propose
from bot.sheet import build_template, parse_workbook
from bot.timeutil import KST, ParseError, parse_time_range, to_epoch

CFG = load_config(ROOT / "config.toml")
G = 1
TODAY = date(2026, 9, 30)  # 수요일 (새 주간 시작)


def kst(*args) -> datetime:
    return datetime(*args, tzinfo=KST)


def fill(content: bytes, sheet: str, rows: list[list]) -> bytes:
    wb = load_workbook(BytesIO(content))
    ws = wb[sheet]
    # 예시 줄은 지우고 채운다
    if str(ws.cell(2, 1).value or "").startswith("예시"):
        ws.delete_rows(2)
    row = 2
    while ws.cell(row, 1).value not in (None, ""):
        row += 1
    for values in rows:  # 사람이 입력하듯 첫 빈 줄부터 채운다 (append 는 서식 칸 아래로 밀림)
        for col, v in enumerate(values, 1):
            ws.cell(row, col, v)
        row += 1
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


class TimeRangeTest(unittest.TestCase):
    def test_ranges(self):
        self.assertEqual(parse_time_range("21-02"), (21 * 60, 26 * 60))
        self.assertEqual(parse_time_range("20:30~24"), (20 * 60 + 30, 24 * 60))
        self.assertEqual(parse_time_range("21~"), (21 * 60, 26 * 60))
        self.assertEqual(parse_time_range("21"), (21 * 60, 26 * 60))
        self.assertEqual(parse_time_range("21시-익일2시"), (21 * 60, 26 * 60))
        self.assertEqual(parse_time_range("14-20"), (14 * 60, 20 * 60))
        with self.assertRaises(ParseError):
            parse_time_range("저녁")


class SheetTest(unittest.TestCase):
    def setUp(self):
        self.db = Database(":memory:")

    def test_empty_template_has_sheets_and_examples_are_ignored(self):
        content = build_template(self.db, CFG, G)
        wb = load_workbook(BytesIO(content))
        self.assertEqual(wb.sheetnames, ["안내", "공대원", "레이드 숙제", "일정", "레이드 이름", "계산용"])
        data = parse_workbook(content, CFG, TODAY)
        self.assertEqual((data.members, data.characters, data.sessions, data.errors), ([], [], [], []))
        self.assertEqual(data.parties, [])  # 일정(파티) 시트는 있지만 예시뿐
        self.assertEqual(wb["계산용"].sheet_state, "hidden")
        self.assertEqual([c.value for c in wb["레이드 이름"][1]], ["레이드", "입장 레벨", "인원"])
        self.assertEqual([c.value for c in wb["공대원"][1]][:2], ["공대원", "대표캐릭터"])
        # 계산 칸이 1000줄까지 깔려 있다 (공대원이 많아도 그대로 이어서 적으면 계산됨)
        self.assertTrue(str(wb["레이드 숙제"]["K1000"].value).startswith("=IF("))
        self.assertNotIn("합계", [c.value for c in wb["레이드 숙제"]["B"]])

    def test_template_prefilled_roundtrip(self):
        m = self.db.upsert_member(G, 123456789012345678, "철수")
        c, _ = self.db.upsert_character(m.id, "철수바드", "바드", 1720.5)
        self.db.set_character_tier(c.id, "1720+_3")
        self.db.upsert_character(m.id, "철수버서", "버서커", 1712)
        self.db.set_availability(m.id, 0, 21 * 60, 26 * 60)
        self.db.set_roster(m.id, "철수바드")
        data = parse_workbook(build_template(self.db, CFG, G, prefill=True), CFG, TODAY)
        self.assertEqual(data.errors, [])
        member = data.members[0]
        self.assertEqual((member.name, member.discord_id, member.roster_name), ("철수", None, "철수바드"))
        self.assertEqual(member.availability, {0: (21 * 60, 26 * 60)})
        self.assertEqual((data.characters[0].name, data.characters[0].item_level), ("철수바드", 1720.5))
        self.assertEqual((data.characters[0].tier, data.characters[1].tier), ("1720+_3", None))  # 두 번째는 '자동'
        self.assertEqual(data.characters[1].member_name, "철수")  # 공대원은 첫 줄에만 적혀 있어도 이어짐
        self.assertEqual(len(data.characters), 2)

    def test_parse_filled_sheet(self):
        content = build_template(self.db, CFG, G)
        content = fill(content, "공대원", [
            ["영희", "영희소서", "21-02", "", "20:30~24", "x", "", "", ""],
            ["지훈", "", "저녁", "", "", "", "", "", ""],            # 잘못된 시간
            ["가람", "", "", "", "", "", "", "", ""],                # 가능 시간 안 건드림
        ])
        content = fill(content, "레이드 숙제", [["가람", "가람건슬", "", "", ""], ["", "가람기상", "기상술사", "1,705.00", "1710+_2"]])
        data = parse_workbook(content, CFG, TODAY)

        names = [m.name for m in data.members]
        self.assertEqual(names, ["영희", "가람"])
        yh = data.members[0]
        self.assertEqual(yh.roster_name, "영희소서")
        self.assertEqual(yh.availability, {0: (1260, 1560), 2: (1230, 1440)})  # '목 x' 는 불가
        self.assertIsNone(data.members[1].availability)
        self.assertEqual(data.characters[0].item_level, None)
        self.assertEqual(data.characters[1].item_level, 1705.0)
        self.assertEqual((data.characters[1].member_name, data.characters[1].tier), ("가람", "1710+_2"))

        errors = "\n".join(data.errors)
        self.assertIn("'저녁'", errors)
        self.assertEqual(data.member_names(), ["영희", "가람"])

    def test_legacy_one_off_schedule_sheet_still_parses(self):
        """이전 양식의 '일정'(날짜·시각·참가 캐릭터) 시트도 읽는다."""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "일정"
        ws.append(["날짜", "시각", "레이드", "참가 캐릭터", "메모"])
        ws.append(["10/2", "21:00", "세르카 노말", "영희소서, 가람건슬", "숙련"])
        ws.append([datetime(2026, 10, 3), time(20, 0), "종막하드", "가람건슬", None])
        out = BytesIO()
        wb.save(out)
        data = parse_workbook(out.getvalue(), CFG, TODAY)
        self.assertIsNone(data.parties)
        self.assertEqual([(s.start, s.raid_key) for s in data.sessions],
                         [(kst(2026, 10, 2, 21), "serka_normal"), (kst(2026, 10, 3, 20), "final_hard")])

    def test_not_an_excel_file(self):
        data = parse_workbook(b"hello", CFG, TODAY)
        self.assertIn("열 수 없어요", data.errors[0])


class PartySheetTest(unittest.TestCase):
    """레이드 파티 시트 → 고정 파티 → 매주 자동 일정 / 시간 미정 파티 제안."""

    def setUp(self):
        from bot.parties import apply_party_inputs
        self.apply = apply_party_inputs
        self.db = Database(":memory:")
        self.names = ["나래", "라온", "다솜", "가온"]
        self.members = {n: self.db.upsert_member(G, 1000 + i, n) for i, n in enumerate(self.names)}
        for n, m in self.members.items():
            self.db.upsert_character(m.id, f"{n}1", "바드", 1785)
            self.db.set_availability(m.id, 4, 20 * 60, 26 * 60)  # 금 20~02

    def party_sheet(self, rows):
        """rows: [파티, 구간, 레이드, 요일, 시각, 캐릭터…] — 캐릭터들은 '캐릭터' 칸 하나에 쉼표로 적는다."""
        content = build_template(self.db, CFG, G)
        wb = load_workbook(BytesIO(content))
        ws = wb["일정"]
        header = [c.value for c in ws[1]]
        self.assertEqual(header, ["파티", "구간", "레이드", "요일", "시각", "캐릭터"])
        ws.delete_rows(2, ws.max_row)  # 예시 줄 지우기
        for r, values in enumerate(rows, 2):
            for col, v in enumerate(values[:5], 1):
                ws.cell(r, col, v)
            ws.cell(r, 6, ", ".join(n for n in values[5:] if n not in ("", "-")))
        out = BytesIO()
        wb.save(out)
        return parse_workbook(out.getvalue(), CFG, TODAY)

    def test_parse_apply_and_generate(self):
        # 머리글 순서: 가온, 나래, 다솜, 라온
        data = self.party_sheet([
            [1, "1780+", "벨가르딘 나이트메어", "금", "21:00", "가온1", "나래1", "다솜1", "라온1"],
            [1, "", "지평의 성당 3단계", "", "", "가온1", "나래1", "다솜1", "라온1"],
            [1, "", "세르카 나이트메어", "금", "", "가온1", "-", "다솜1", "라온1"],   # 요일만 → 기본 시각(21:00)에 그 요일로
            [2, "", "세르카 노말", "토", "20시", "", "없는캐릭", "", ""],
        ])
        self.assertEqual([p.number for p in data.parties], [1, 2])
        p1 = data.parties[0]
        self.assertEqual((p1.label, len(p1.rows)), ("1780+", 3))
        self.assertEqual((p1.rows[0].weekday, p1.rows[0].start_min), (4, 21 * 60))
        self.assertIsNone(p1.rows[1].weekday)
        self.assertEqual(len(p1.rows[2].entries), 3)  # '-' 는 빠짐
        self.assertEqual((p1.rows[2].weekday, p1.rows[2].start_min), (4, 21 * 60))  # 시각 비움 → 기본 시각
        self.assertFalse(any("요일" in e for e in data.errors))

        ids = {n: m.id for n, m in self.members.items()}
        errors = self.apply(self.db, G, data.parties, ids)
        self.assertTrue(any("없는캐릭" in e for e in errors))
        parties = self.db.list_parties(G)
        self.assertEqual(len(parties[0].raids[0].characters), 4)

        from bot.parties import generate_party_sessions
        now = kst(2026, 9, 30, 12)
        gen = generate_party_sessions(self.db, CFG, G, now)
        # 파티1 벨가(금 21시) + 파티1 세르카 나메(금, 시각 비움=기본 21시, 벨가와 겹쳐도 허용) + 파티2 세르카 노말(토 20시)
        self.assertEqual(len(gen.created), 3)
        # 두 번 불러도 같은 주에는 다시 만들지 않는다
        self.assertEqual(generate_party_sessions(self.db, CFG, G, now).created, [])
        sessions = self.db.list_sessions(G, to_epoch(kst(2026, 9, 30, 6)), to_epoch(kst(2026, 10, 7, 6)))
        self.assertEqual({(s.raid_key, s.start_at) for s in sessions},
                         {("belgardin_nightmare", to_epoch(kst(2026, 10, 2, 21))),
                          ("serka_nightmare", to_epoch(kst(2026, 10, 2, 21))),
                          ("serka_normal", to_epoch(kst(2026, 10, 3, 20)))})
        # 같은 시각에 두 일정이 겹쳐도 같은 파티원이 양쪽에 다 들어간다
        both = [len(self.db.session_characters(x.id)) for x in sessions if x.start_at == to_epoch(kst(2026, 10, 2, 21))]
        self.assertEqual(both, [4, 3] if both[0] == 4 else [3, 4])

        # 요일도 시각도 없는 줄(성당 3단계)만 파티 단위로 제안. 자동 제안은 겹치지 않게 벨가·세르카(21~) 뒤에 붙는다
        result = propose(self.db, CFG, G, now)
        party_props = [p for p in result.proposals if p.party_number == 1]
        self.assertEqual([p.raid.key for p in party_props], ["cathedral_3"])
        self.assertFalse(party_props[0].start < kst(2026, 10, 2, 22) and party_props[0].end > kst(2026, 10, 2, 21))
        # 파티가 맡은 레이드는 일반 매칭에 다시 나오지 않는다
        self.assertFalse(any(p.party_number is None and p.raid.key == "cathedral_3" for p in result.proposals))

    def test_time_without_day_is_flagged(self):
        data = self.party_sheet([[3, "1780+", "세르카 나이트메어", "", "22:00", "가온1", "나래1", "", ""]])
        self.assertTrue(any("요일이 없어요" in e for e in data.errors))
        self.assertIsNone(data.parties[0].rows[0].weekday)

    def test_template_prefills_parties_with_colors(self):
        ids = {n: m.id for n, m in self.members.items()}
        data = self.party_sheet([[3, "1780+", "세르카 나이트메어", "수", "22:00", "가온1", "나래1", "", ""]])
        self.apply(self.db, G, data.parties, ids)
        wb = load_workbook(BytesIO(build_template(self.db, CFG, G, prefill=True)))
        ws = wb["일정"]
        self.assertEqual([ws.cell(2, c).value for c in range(1, 7)],
                         [3, "1780+", "세르카 나이트메어", "수", "22:00", "가온1, 나래1"])
        self.assertGreaterEqual(len(ws.conditional_formatting), 1)  # 파티 번호별 색 규칙


    def test_default_template_has_only_example_data(self):
        """기본 양식은 등록된 공대원 이름 없이 예시(로디코 / 로워붕·로소서·로권왕)만. 예시 줄은 올려도 무시."""
        content = build_template(self.db, CFG, G)
        wb = load_workbook(BytesIO(content))
        values = {str(c.value) for ws in wb.worksheets if ws.title != "계산용"
                  for row in ws.iter_rows(max_row=10) for c in row if c.value is not None}
        for real in self.names:
            self.assertFalse(any(real in v for v in values), real)
        self.assertIn("로디코", values)
        self.assertTrue({"로워붕", "로소서", "로권왕"} <= values)
        data = parse_workbook(content, CFG, TODAY)
        self.assertEqual((data.members, data.characters, data.parties), ([], [], []))

    def test_legacy_member_columns_still_parse(self):
        """이전 양식(공대원 이름 열마다 캐릭터)도 읽는다."""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "일정"
        ws.append(["파티", "구간", "레이드", "요일", "시각", "가온", "나래", "비고"])
        ws.append([1, "", "세르카 나이트메어", "금", "21:00", "가온1", "나래1", "메모"])
        out = BytesIO()
        wb.save(out)
        data = parse_workbook(out.getvalue(), CFG, TODAY)
        self.assertEqual(data.parties[0].rows[0].entries, [("가온", "가온1"), ("나래", "나래1")])


class AvailableBetweenTest(unittest.TestCase):
    def test_windows(self):
        db = Database(":memory:")
        m = db.upsert_member(G, 1, "a")
        db.set_availability(m.id, 4, 21 * 60, 26 * 60)  # 금 21 ~ 토 02
        slots = db.list_availability(m.id)
        self.assertTrue(available_between(slots, kst(2026, 10, 2, 21), kst(2026, 10, 2, 22)))
        self.assertTrue(available_between(slots, kst(2026, 10, 3, 1), kst(2026, 10, 3, 2)))   # 전날 창 이어짐
        self.assertFalse(available_between(slots, kst(2026, 10, 3, 1, 30), kst(2026, 10, 3, 2, 30)))
        self.assertFalse(available_between(slots, kst(2026, 10, 2, 20), kst(2026, 10, 2, 21, 30)))


class PlannerTest(unittest.TestCase):
    def setUp(self):
        self.db = Database(":memory:")
        self.now = kst(2026, 9, 30, 12)  # 수 12시

    def add(self, name: str, levels: list[float], avail: dict[int, tuple[int, int]] | None, cls: str | None = None):
        # 첫 공대원(p0)과 8인용 두 번째(p1)는 서폿(바드), 나머지는 딜러
        idx = len(self.db.list_members(G))
        cls = cls or ("바드" if idx == 0 else "블레이드")
        m = self.db.upsert_member(G, 1000 + idx, name)
        chars = [self.db.upsert_character(m.id, f"{name}{i}", cls, lv)[0] for i, lv in enumerate(levels)]
        for d, (s, e) in (avail or {}).items():
            self.db.set_availability(m.id, d, s, e)
        return m, chars

    def test_fills_4_person_raid_in_common_window(self):
        # 4명 모두 1715 (세르카 노말/종막 노말/성당 1단계가 골드 레이드), 금 21~24 만 공통
        for i in range(4):
            self.add(f"p{i}", [1715], {4: (21 * 60, 24 * 60), i % 3: (10 * 60, 11 * 60)})
        result = propose(self.db, CFG, G, self.now, raid_key="serka_normal")
        self.assertEqual(len(result.proposals), 1)
        p = result.proposals[0]
        self.assertEqual(p.raid.key, "serka_normal")
        self.assertEqual(p.start, kst(2026, 10, 2, 21))
        self.assertEqual(len({c.member_id for c in p.characters}), 4)
        self.assertEqual(result.leftovers, {})

    def test_not_enough_people_goes_to_leftovers_or_partial(self):
        for i in range(3):
            self.add(f"p{i}", [1715], {4: (21 * 60, 24 * 60)})
        result = propose(self.db, CFG, G, self.now, raid_key="serka_normal")
        self.assertEqual(result.proposals, [])
        self.assertEqual(len(result.leftovers["serka_normal"]), 3)
        partial = propose(self.db, CFG, G, self.now, raid_key="serka_normal", min_fill=3)
        self.assertEqual(len(partial.proposals[0].characters), 3)

    def test_member_with_two_characters_not_double_booked(self):
        # 4명 각자 캐릭터 2개가 세르카 노말 필요 → 세션 2개, 각 세션에 공대원당 1캐릭, 시간 안 겹침
        for i in range(4):
            self.add(f"p{i}", [1715, 1712], {4: (20 * 60, 24 * 60)})
        result = propose(self.db, CFG, G, self.now, raid_key="serka_normal")
        self.assertEqual(len(result.proposals), 2)
        a, b = result.proposals
        self.assertTrue(a.end <= b.start or b.end <= a.start)
        for p in result.proposals:
            self.assertEqual(len({c.member_id for c in p.characters}), 4)
        used = [c.id for p in result.proposals for c in p.characters]
        self.assertEqual(len(used), len(set(used)))

    def test_existing_sessions_block_time_and_demand(self):
        members = [self.add(f"p{i}", [1715], {4: (21 * 60, 23 * 60)}) for i in range(4)]
        # p0 는 금 21시에 다른 일정이 있고, 이미 세르카 노말도 잡혀 있지 않음
        other = self.db.create_session(G, "final_normal", to_epoch(kst(2026, 10, 2, 21)), 1, None)
        self.db.add_slot(other.id, members[0][1][0].id)
        result = propose(self.db, CFG, G, self.now, raid_key="serka_normal")
        # p0 가 21~22 에 바쁘므로 22시 시작으로 4명 채움
        self.assertEqual(result.proposals[0].start, kst(2026, 10, 2, 22))

    def test_big_party_is_placed_before_small_ones_take_the_slot(self):
        # 8명 모두 금 21~23시만 가능. 골드 레이드 = 세르카 나메(4인), 종막 하드(8인), 성당 2단계(4인)
        for i in range(8):
            self.add(f"p{i}", [1745], {4: (21 * 60, 23 * 60)}, cls="바드" if i < 2 else None)
        result = propose(self.db, CFG, G, self.now)
        final = [p for p in result.proposals if p.raid.key == "final_hard"]
        self.assertEqual(len(final), 1)
        self.assertEqual(len(final[0].characters), 8)
        # 남은 한 시간(22시)에 4인 두 파티
        others = [p for p in result.proposals if p.raid.party_size == 4]
        self.assertEqual({p.start for p in others}, {kst(2026, 10, 2, 22)})
        self.assertEqual(len(others), 2)

    def test_support_rule_in_proposals(self):
        # 딜러 5명만 있으면 4인 파티를 못 채운다 (딜러 최대 3명)
        for i in range(5):
            self.add(f"d{i}", [1715], {4: (21 * 60, 24 * 60)}, cls="블레이드")
        self.assertEqual(propose(self.db, CFG, G, self.now, raid_key="serka_normal").proposals, [])
        # 서폿이 생기면 서폿 1 + 딜러 3
        self.add("s0", [1715], {4: (21 * 60, 24 * 60)}, cls="홀리나이트")
        (p,) = propose(self.db, CFG, G, self.now, raid_key="serka_normal").proposals
        self.assertEqual(sum(CFG.is_support(c.cls) for c in p.characters), 1)
        self.assertEqual(len(p.characters), 4)

    def test_no_availability_listed(self):
        self.add("유령", [1715], None)
        result = propose(self.db, CFG, G, self.now)
        self.assertEqual(result.no_availability, ["유령"])
        self.assertEqual(result.proposals, [])

    def test_does_not_propose_in_the_past_or_next_week(self):
        for i in range(4):
            self.add(f"p{i}", [1715], {d: (10 * 60, 11 * 60) for d in range(7)})
        result = propose(self.db, CFG, G, kst(2026, 9, 30, 10, 30), raid_key="serka_normal")
        for p in result.proposals:
            self.assertGreaterEqual(p.start, kst(2026, 9, 30, 11, 30))
            self.assertLess(p.end, kst(2026, 10, 7, 6) + timedelta(seconds=1))


if __name__ == "__main__":
    unittest.main()


class UnavailableTest(unittest.TestCase):
    """/불가 예외: 자동 제안에서 제외, 참가 경고, 만료 정리."""
    setUp = PlannerTest.setUp
    add = PlannerTest.add

    def test_exception_moves_proposal_to_other_time(self):
        members = [self.add(f"p{i}", [1715], {4: (21 * 60, 24 * 60), 5: (21 * 60, 24 * 60)}) for i in range(4)]
        # 금 21~24 에 p0 가 불가 → 토요일로 밀림
        self.db.add_unavailable(members[0][0].id, to_epoch(kst(2026, 10, 2, 21)), to_epoch(kst(2026, 10, 3, 0)))
        result = propose(self.db, CFG, G, self.now, raid_key="serka_normal")
        self.assertEqual(result.proposals[0].start, kst(2026, 10, 3, 21))

    def test_all_day_exception_and_purge(self):
        m = self.db.upsert_member(G, 1, "a")
        start, end = build_day(2026, 10, 2)
        u = self.db.add_unavailable(m.id, start, end)
        self.assertEqual([x.id for x in self.db.list_unavailable(m.id)], [u.id])
        self.assertEqual(self.db.purge_unavailable(start), 0)
        self.assertEqual(self.db.purge_unavailable(end), 1)
        self.assertEqual(self.db.list_unavailable(m.id), [])

    def test_delete_only_own(self):
        a, b = self.db.upsert_member(G, 1, "a"), self.db.upsert_member(G, 2, "b")
        u = self.db.add_unavailable(a.id, 100, 200)
        self.assertFalse(self.db.delete_unavailable(b.id, u.id))
        self.assertTrue(self.db.delete_unavailable(a.id, u.id))

    def test_join_warns_when_overlapping(self):
        from bot.services import check_join
        m, chars = self.add("p", [1715], {4: (21 * 60, 24 * 60)})
        s = self.db.create_session(G, "serka_normal", to_epoch(kst(2026, 10, 2, 21)), 1, None)
        self.assertFalse(any("불가" in w for w in check_join(self.db, CFG, s, chars[0], self.now)))
        self.db.add_unavailable(m.id, to_epoch(kst(2026, 10, 2, 20)), to_epoch(kst(2026, 10, 2, 22)))
        self.assertTrue(any("불가" in w for w in check_join(self.db, CFG, s, chars[0], self.now)))

    def test_build_interval(self):
        from bot.cogs.unavailable import build_interval
        today = date(2026, 9, 30)
        self.assertEqual(build_interval("10/3", "20~24", today),
                         (to_epoch(kst(2026, 10, 3, 20)), to_epoch(kst(2026, 10, 4, 0))))
        self.assertEqual(build_interval("10/3", None, today), build_day(2026, 10, 3))
        self.assertEqual(build_interval("10/3", "21~02", today)[1], to_epoch(kst(2026, 10, 4, 2)))


def build_day(y: int, mo: int, d: int) -> tuple[int, int]:
    return to_epoch(kst(y, mo, d)), to_epoch(kst(y, mo, d) + timedelta(days=1))


class ResetEveTest(unittest.TestCase):
    def test_summary_lists_only_unfinished_raids(self):
        from bot.render import reset_eve_text
        from bot.services import build_overview
        db = Database(":memory:")
        m = db.upsert_member(G, 111, "철수")
        c = db.upsert_character(m.id, "철수1", "블레이드", 1715)[0]
        now = kst(2026, 10, 6, 9)  # 화 09시, 다음 날(수) 06시 리셋
        text = reset_eve_text(CFG, build_overview(db, CFG, G, now), now)
        self.assertIn("<@111>", text)
        self.assertIn("❔", text)
        # 골드 레이드 전부를 이미 끝난 일정에 넣으면 요약할 게 없다
        for r in [x for x in build_overview(db, CFG, G, now).members[0].characters[0].gold_raids]:
            s = db.create_session(G, r.key, to_epoch(kst(2026, 10, 3, 20)), 1, None)
            db.add_slot(s.id, c.id)
        self.assertIsNone(reset_eve_text(CFG, build_overview(db, CFG, G, now), now))
        # 예정(아직 안 끝난) 일정 하나만 있는 상태는 🗓️ 로 표시
        db2 = Database(":memory:")
        m2 = db2.upsert_member(G, 222, "영희")
        c2 = db2.upsert_character(m2.id, "영희1", "블레이드", 1715)[0]
        s2 = db2.create_session(G, "serka_normal", to_epoch(kst(2026, 10, 6, 21)), 1, None)
        db2.add_slot(s2.id, c2.id)
        text2 = reset_eve_text(CFG, build_overview(db2, CFG, G, now), now)
        self.assertIn("🗓️", text2)
        self.assertIn("❔", text2)  # 나머지 골드 레이드는 아직 안 잡힘


class GoogleExportTest(unittest.TestCase):
    def test_build_export_tables(self):
        from bot.gsheet import TAB_CLEARS, TAB_SESSIONS, TAB_STATS, build_export, load_export_config
        db = Database(":memory:")
        m = db.upsert_member(G, 1, "철수")
        c = db.upsert_character(m.id, "철수1", "블레이드", 1715)[0]
        now = kst(2026, 10, 6, 12)
        past = db.create_session(G, "serka_normal", to_epoch(kst(2026, 10, 3, 20)), 1, None)
        db.add_slot(past.id, c.id)
        future = db.create_session(G, "final_normal", to_epoch(kst(2026, 10, 6, 21)), 1, None)
        db.add_slot(future.id, c.id)
        t = build_export(db, CFG, G, now)
        sess = t[TAB_SESSIONS]
        self.assertEqual(len(sess), 3)
        self.assertEqual([r[5] for r in sess[1:]], ["✅ 완료", "🗓️ 예정"])
        statuses = {r[4]: r[5] for r in t[TAB_CLEARS][1:]}
        self.assertEqual(sum(1 for s in statuses.values() if s == "✅ 완료"), 1)
        self.assertTrue(any(s.startswith("❔") for s in statuses.values()))
        row = t[TAB_STATS][1]
        self.assertEqual(row[0], "철수")
        self.assertEqual(row[6], "1")          # 끝난 일정 1개
        self.assertEqual(int(row[1]), int(row[2]) + int(row[3]) + int(row[4]))
        self.assertTrue(all(isinstance(x, str) for r in sum(t.values(), []) for x in r))
        self.assertIsNone(load_export_config(None, "abc"))
        self.assertIsNotNone(load_export_config("k.json", "abc"))

    def test_missing_key_file_gives_friendly_error(self):
        from bot.gsheet import ExportConfig, ExportError, write_export
        with self.assertRaises(ExportError):
            write_export(ExportConfig("no-such-key.json", "id"), {"a": [["x"]]})


class SessionButtonTest(unittest.TestCase):
    def setUp(self):
        self.db = Database(":memory:")

        class B:  # join_session 이 쓰는 최소한의 봇
            pass
        self.bot = B()
        self.bot.db, self.bot.cfg = self.db, CFG

    def test_join_and_leave_texts(self):
        from bot.cogs.sessions import join_session, leave_session
        from bot.services import RuleError
        m = self.db.upsert_member(G, 1, "철수")
        c = self.db.upsert_character(m.id, "철수1", "블레이드", 1715)[0]
        s = self.db.create_session(G, "serka_normal", to_epoch(datetime.now(KST) + timedelta(days=1)), 1, None)
        self.assertIn("철수1", join_session(self.bot, s, self.db.get_character(c.id), by_leader=False))
        with self.assertRaises(RuleError):  # 이미 참가 중
            join_session(self.bot, s, self.db.get_character(c.id), by_leader=False)
        self.assertIn("빠졌어요", leave_session(self.bot, s, m))
        with self.assertRaises(RuleError):
            leave_session(self.bot, s, m)

    def test_custom_id_pattern_matches(self):
        from bot.cogs.sessions import _SessionButton
        self.assertTrue(_SessionButton.__discord_ui_compiled_template__.fullmatch("raid:join:12"))
        self.assertTrue(_SessionButton.__discord_ui_compiled_template__.fullmatch("raid:leave:3"))
        self.assertIsNone(_SessionButton.__discord_ui_compiled_template__.fullmatch("raid:x:3"))


class DevSeedTest(unittest.TestCase):
    def test_add_propose_remove_leaves_real_members_alone(self):
        from bot.devseed import add_dummies, list_dummies, remove_dummies
        db = Database(":memory:")
        real = db.upsert_member(G, 123456789012345678, "진짜")
        db.upsert_character(real.id, "진짜1", "블레이드", 1720)
        lines = add_dummies(db, G, 8)
        self.assertEqual(len(lines), 8)
        self.assertEqual(len(list_dummies(db, G)), 8)
        result = propose(db, CFG, G, kst(2026, 9, 30, 12))
        self.assertGreater(len(result.proposals), 0)      # 더미만으로도 매칭이 돌아간다
        self.assertEqual(remove_dummies(db, G), 8)
        self.assertEqual([m.display_name for m in db.list_members(G)], ["진짜"])
        self.assertEqual(len(db.list_characters(real.id)), 1)


class SessionEditTest(unittest.TestCase):
    def setUp(self):
        from bot.session_edit import EditRequest
        self.R = EditRequest
        self.db = Database(":memory:")
        self.now = kst(2026, 9, 30, 12)
        self.a = self.db.upsert_member(G, 1, "가")
        self.b = self.db.upsert_member(G, 2, "나")
        self.ca = self.db.upsert_character(self.a.id, "가1", "블레이드", 1720)[0]
        self.ca2 = self.db.upsert_character(self.a.id, "가2", "소서리스", 1715)[0]
        self.cb = self.db.upsert_character(self.b.id, "나1", "워로드", 1720)[0]
        self.s1 = self.db.create_session(G, "serka_normal", to_epoch(kst(2026, 10, 2, 21)), 1, None)
        self.s2 = self.db.create_session(G, "final_normal", to_epoch(kst(2026, 10, 3, 21)), 1, "⏱️ 시간 미정 (임시 21:00) · 숙련")
        self.s3 = self.db.create_session(G, "serka_hard", to_epoch(kst(2026, 10, 4, 21)), 1, None)
        self.db.add_slot(self.s1.id, self.ca.id)
        self.db.add_slot(self.s2.id, self.ca.id)
        self.db.add_slot(self.s3.id, self.cb.id)

    def edit(self, sid, req, member=None, leader=False):
        from bot.session_edit import apply_edit
        return apply_edit(self.db, CFG, G, sid, req, (member or self.a).id, leader, self.now)

    def test_list_respects_permission(self):
        from bot.session_edit import editable_sessions
        mine = editable_sessions(self.db, CFG, G, self.now, self.a.id, False)
        self.assertEqual([v.session.id for v in mine], [self.s1.id, self.s2.id])
        everyone = editable_sessions(self.db, CFG, G, self.now, self.a.id, True)
        self.assertEqual(len(everyone), 3)

    def test_member_cannot_edit_others_session_but_leader_can(self):
        from bot.services import RuleError
        with self.assertRaisesRegex(RuleError, "내가 참가한 일정이 아니라"):
            self.edit(self.s3.id, self.R(time_text="22:00"))
        self.edit(self.s3.id, self.R(time_text="22:00"), leader=True)
        self.assertEqual(self.db.get_session(self.s3.id).start_at, to_epoch(kst(2026, 10, 4, 22)))

    def test_shift_and_date_time(self):
        self.edit(self.s1.id, self.R(shift_minutes=60))
        self.assertEqual(self.db.get_session(self.s1.id).start_at, to_epoch(kst(2026, 10, 2, 22)))
        self.edit(self.s1.id, self.R(date_text="10/5", time_text="20:00", shift_minutes=-30))
        self.assertEqual(self.db.get_session(self.s1.id).start_at, to_epoch(kst(2026, 10, 5, 19, 30)))

    def test_time_tbd_marker_removed_when_time_set(self):
        self.edit(self.s2.id, self.R(time_text="22:00"))
        self.assertEqual(self.db.get_session(self.s2.id).memo, "숙련")
        # 메모를 직접 고치면 그 메모가 우선
        self.edit(self.s2.id, self.R(memo="-"))
        self.assertIsNone(self.db.get_session(self.s2.id).memo)

    def test_validation(self):
        from bot.services import RuleError
        from bot.session_edit import parse_shift
        with self.assertRaises(RuleError):
            self.edit(self.s1.id, self.R())                       # 비어 있음
        with self.assertRaisesRegex(RuleError, "지난 시각"):
            self.edit(self.s1.id, self.R(date_text="9/29"))       # 과거
        with self.assertRaises(RuleError):
            parse_shift("한시간")
        self.assertEqual((parse_shift("+60"), parse_shift("-30"), parse_shift("")), (60, -30, 0))

    def test_leave_and_swap(self):
        from bot.services import RuleError
        from bot.session_edit import leave_mine, swap_mine
        swap_mine(self.db, CFG, G, self.s1.id, self.db.get_character(self.ca2.id), self.now)
        self.assertEqual([c.name for c in self.db.session_characters(self.s1.id)], ["가2"])
        # 같은 레이드 주 1회 위반이면 교체가 실패하고 원래대로 돌아온다 (가1 은 s2 의 final_normal 에 이미 있음)
        s_final = self.db.create_session(G, "final_hard", to_epoch(kst(2026, 10, 4, 22)), 1, None)
        self.db.add_slot(s_final.id, self.ca2.id)
        with self.assertRaises(RuleError):
            swap_mine(self.db, CFG, G, s_final.id, self.db.get_character(self.ca.id), self.now)
        self.assertEqual([c.name for c in self.db.session_characters(s_final.id)], ["가2"])
        leave_mine(self.db, CFG, G, self.s1.id, self.a.id, self.now)
        self.assertEqual(self.db.session_characters(self.s1.id), [])
        with self.assertRaises(RuleError):
            leave_mine(self.db, CFG, G, self.s1.id, self.a.id, self.now)


class SessionEditViewTest(unittest.IsolatedAsyncioTestCase):
    async def test_view_builds_for_member_leader_and_empty(self):
        from types import SimpleNamespace
        from bot.cogs.session_edit import SessionEditView
        db = Database(":memory:")
        a = db.upsert_member(G, 1, "가")
        c = db.upsert_character(a.id, "가1", "블레이드", 1720)[0]
        start = to_epoch(datetime.now(KST) + timedelta(days=2))
        s = db.create_session(G, "serka_normal", start, 1, None)
        db.add_slot(s.id, c.id)
        bot = SimpleNamespace(db=db, cfg=CFG)
        mine = SessionEditView(bot, G, 1, a.id, False)
        self.assertEqual(len(mine.sessions), 1)
        self.assertIn(f"#{s.id}", mine.content())
        self.assertEqual(len(mine.picker.options), 1)
        other = db.upsert_member(G, 2, "나")
        none = SessionEditView(bot, G, 2, other.id, False)       # 참가한 일정이 없는 사람
        self.assertEqual(none.sessions, [])
        self.assertIn("없어요", none.content())
        self.assertFalse(any(getattr(i, "label", "") == "시각·메모 수정" for i in none.children))
        leader = SessionEditView(bot, G, 2, other.id, True)      # 공대장은 남의 일정도 보인다
        self.assertEqual(len(leader.sessions), 1)


class ClearedAndColorTest(unittest.TestCase):
    def setUp(self):
        self.db = Database(":memory:")
        self.now = kst(2026, 10, 3, 12)      # 토 12시 (주간: 9/30 06:00 ~ 10/7 06:00)
        self.m = self.db.upsert_member(G, 1, "가")
        self.other = self.db.upsert_member(G, 2, "나")
        self.c = self.db.upsert_character(self.m.id, "가1", "블레이드", 1725)[0]
        self.c2 = self.db.upsert_character(self.other.id, "나1", "워로드", 1725)[0]

    def make(self, raid, when, memo=None, party_raid_id=None, chars=()):
        s = self.db.create_session(G, raid, to_epoch(when), 1, memo, party_raid_id=party_raid_id)
        for c in chars:
            self.db.add_slot(s.id, c.id)
        return s

    def test_party_color_falls_back_to_memo(self):
        from bot.render import overview_text
        from bot.services import build_overview
        old = self.make("serka_normal", kst(2026, 10, 4, 21), memo="파티 3", chars=[self.c])   # 파티 연결 없이 만든 예전 일정
        plain = self.make("final_normal", kst(2026, 10, 4, 22), memo="숙련", chars=[self.c])
        ov = build_overview(self.db, CFG, G, self.now)
        self.assertEqual(ov.party_of.get(old.id), 3)
        self.assertNotIn(plain.id, ov.party_of)
        text = overview_text(CFG, ov, self.now)
        self.assertIn("🟧P3", text)           # 3번 파티 색 + 번호
        self.assertNotIn("· 파티 3", text)    # 색이 있으니 '파티 N' 메모는 생략
        self.assertIn("▫️ **22:00", text)     # 파티가 아닌 일정도 같은 자리에 표시
        self.assertIn("· 숙련", text)         # 다른 메모는 그대로

    def test_manual_done_and_failed(self):
        from bot.render import overview_text, session_done
        from bot.services import SessionView, build_overview
        from bot.session_edit import mark_cleared
        from bot.services import RuleError
        future = self.make("serka_normal", kst(2026, 10, 5, 21), chars=[self.c])
        past = self.make("final_normal", kst(2026, 10, 2, 21), chars=[self.c])
        v = lambda s: SessionView(self.db.get_session(s.id), CFG.raid(s.raid_key), [self.c])
        self.assertTrue(session_done(v(past), self.now))               # 끝 시각이 지나 자동 완료
        with self.assertRaisesRegex(RuleError, "시작 전"):
            mark_cleared(self.db, G, future.id, 1, self.m.id, False, self.now)
        with self.assertRaisesRegex(RuleError, "내가 참가한 일정이 아니라"):
            mark_cleared(self.db, G, past.id, 0, self.other.id, False, self.now)
        mark_cleared(self.db, G, past.id, 0, self.m.id, False, self.now)  # 참가자는 표시 가능
        self.assertFalse(session_done(v(past), self.now))
        text = overview_text(CFG, build_overview(self.db, CFG, G, self.now), self.now)
        self.assertIn("❌", text)
        self.assertIn("**미완료**", text)
        # 미완료로 표시한 레이드는 다시 '안 잡힌 숙제'로 돌아온다
        ov = build_overview(self.db, CFG, G, self.now)
        self.assertIn(self.c.id, [c.id for c in ov.unassigned_by_raid.get("final_normal", [])])
        mark_cleared(self.db, G, past.id, None, None, True, self.now)      # 공대장이 자동으로 되돌림
        self.assertTrue(session_done(v(past), self.now))
        mark_cleared(self.db, G, past.id, 1, self.m.id, False, self.now)   # 수동 완료
        ov = build_overview(self.db, CFG, G, self.now)
        self.assertNotIn(self.c.id, [c.id for c in ov.unassigned_by_raid.get("final_normal", [])])

    def test_failed_session_lets_same_raid_be_rejoined_and_reproposed(self):
        from bot.services import check_join
        past = self.make("serka_normal", kst(2026, 10, 2, 21), chars=[self.c])
        again = self.make("serka_normal", kst(2026, 10, 6, 21))
        from bot.services import RuleError
        with self.assertRaisesRegex(RuleError, "이미 배정"):
            check_join(self.db, CFG, again, self.c, self.now)             # 완료(자동)한 레이드는 주 1회 규칙
        self.db.set_cleared(past.id, 0)
        check_join(self.db, CFG, again, self.c, self.now)                 # 미완료면 다시 참가 가능

    def test_panel_pages_and_marks(self):
        import asyncio
        from types import SimpleNamespace
        from bot.cogs.session_edit import PAGE_SIZE, SessionEditView

        async def run():
            start = datetime.now(KST) + timedelta(hours=2)
            for i in range(PAGE_SIZE + 5):
                s = self.db.create_session(G, "serka_normal", to_epoch(start + timedelta(hours=i)), 1, None)
                self.db.add_slot(s.id, self.c.id)
            bot = SimpleNamespace(db=self.db, cfg=CFG)
            view = SessionEditView(bot, G, 1, self.m.id, False)
            self.assertEqual(view.pages, 2)
            self.assertEqual(len(view.picker.options), PAGE_SIZE)
            view.page = 1
            view._rebuild()
            self.assertEqual(len(view.picker.options), 5)
            self.assertFalse(view.prev_button.disabled)
            self.assertTrue(view.next_button.disabled)
        asyncio.run(run())


class PartyRelinkTest(unittest.TestCase):
    def test_existing_party_session_is_linked_not_duplicated(self):
        from bot.parties import generate_party_sessions
        db = Database(":memory:")
        m = db.upsert_member(G, 1, "가")
        c = db.upsert_character(m.id, "가1", "바드", 1785)[0]
        db.replace_parties(G, [{"number": 1, "label": "1780+", "raids": [
            {"raid_key": "belgardin_nightmare", "weekday": 4, "start_min": 21 * 60, "character_ids": [c.id]}]}])
        # 파티 연결 없이 예전에 다른 시각으로 만든 같은 파티·레이드 일정
        old = db.create_session(G, "belgardin_nightmare", to_epoch(kst(2026, 10, 2, 22)), 1, "파티 1")
        now = kst(2026, 9, 30, 12)
        gen = generate_party_sessions(db, CFG, G, now)
        self.assertEqual(gen.created, [])                       # 새로 만들지 않음
        sessions = db.list_sessions(G, to_epoch(kst(2026, 9, 30, 6)), to_epoch(kst(2026, 10, 7, 6)))
        self.assertEqual([s.id for s in sessions], [old.id])
        self.assertIsNotNone(db.get_session(old.id).party_raid_id)   # 파티에 연결됨
        self.assertEqual(generate_party_sessions(db, CFG, G, now).created, [])

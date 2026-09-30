"""공대 입력용 엑셀(.xlsx) 양식 만들기와 읽기. 디스코드와 무관한 순수 로직.

공대에서 쓰던 시트(레이드 숙제 / 레이드 파티)를 바탕으로 한 구성.

- 안내      : 작성 방법
- 공대원    : 공대원 | 대표캐릭터 | 월 … 일 (가능 시간)
- 레이드 숙제: 공대원 | 캐릭터 | 직업 | 아이템레벨 | 숙제 구간 | (자동 계산) 적용 구간 · 레이드1~3 · 골드X · 클골 · 더보기 · 더보기 후
- 일정      : 파티 | 구간 | 레이드 | 요일 | 시각 | 캐릭터 (여러 명은 쉼표로)   (고정 파티, 파티 번호별 색 자동)
- 레이드 이름: 레이드 | 입장 레벨 | 인원 — 드롭다운 원본
- 계산용    : (숨김) 숙제 구간표 · 레이드 골드/더보기 — 레이드 숙제 수식이 참조
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from .config import EVENT_PREFIX, Config
from .db import Database
from .timeutil import KST, WEEKDAYS, ParseError, fmt_minutes, fmt_range, parse_date, parse_time, parse_time_range

SHEET_GUIDE, SHEET_MEMBERS, SHEET_HOMEWORK = "안내", "공대원", "레이드 숙제"
SHEET_SCHEDULE, SHEET_RAID_NAMES, SHEET_DATA = "일정", "레이드 이름", "계산용"
SHEET_PARTY_OLD, SHEET_CHARS_OLD = "레이드 파티", "캐릭터"  # 이전 양식 호환

MEMBER_HEADERS = ["공대원", "대표캐릭터", *WEEKDAYS]
HOMEWORK_HEADERS = ["공대원", "캐릭터", "직업", "아이템레벨", "숙제 구간"]
HOMEWORK_CALC = ["적용 구간", "레이드1", "레이드2", "레이드3", "골드X", "클골", "더보기 비용", "더보기 후"]
PARTY_FIXED = ["파티", "구간", "레이드", "요일", "시각"]
PARTY_CHARS = "캐릭터"          # 같이 가는 캐릭터 이름들 (쉼표로 구분)
LEGACY_NOTE = {"비고", "메모"}  # 이전 양식의 메모 열 (공대원 열이 아님)
SESSION_HEADERS = ["날짜", "시각", "레이드", "참가 캐릭터", "메모"]  # 이전 양식 '일정' (한 번만 잡는 일정)
EXAMPLE_PREFIX = "예시"
TOTAL_ROW = "합계"  # 이전 양식 호환 (읽을 때 건너뜀)
AUTO = "자동"
MAX_ROWS = 1000     # 수식·드롭다운·서식을 미리 깔아 두는 줄 수 (공대원이 많아도 이만큼은 그대로 적으면 됨)
SKIP = {"", "-", "x", "X"}


# 파티 번호별 배경색 (parties.PALETTE 와 같은 순서)
PARTY_FILLS = ["D9D2E9", "CFE2F3", "FBE4D5", "E2EFD9", "FEF2CB", "F4CCCC", "E6D5C3", "EDEDED"]

# 예시 데이터 (회색 글씨). 이 이름들로 된 줄은 올려도 무시한다.
EX_MEMBER = "로디코"
EX_CHARS = [("로워붕", "워로드", "1780", AUTO), ("로소서", "소서리스", "1752", AUTO), ("로권왕", "브레이커", "1735", AUTO)]
EX_CHAR_NAMES = {c[0] for c in EX_CHARS}
EX_SCHEDULE = [  # 파티, 구간, 레이드, 요일, 시각, 캐릭터 (같이 가는 캐릭터를 쉼표로 여러 명)
    (1, "1730+", "종막 하드", "수", "21:00", "로워붕, 로소서, 로권왕"),
    (1, "1730+", "지평의 성당 2단계", "수", "22:00", "로워붕, 로소서, 로권왕"),
    (1, "1730+", "세르카 하드", "목", "21:00", "로워붕, 로소서, 로권왕"),
    (2, "1750+", "벨가르딘 노말", "금", "21:00", "로워붕, 로소서"),
    (2, "1750+", "세르카 나이트메어", "금", "22:00", "로워붕, 로소서"),
]

GUIDE = [
    "공대 일정 봇 입력 양식",
    "",
    "[공대원] 한 줄에 한 명. 이름은 디스코드 서버에서 보이는 이름(별명)과 똑같이.",
    "   · 대표캐릭터를 적으면 공식 API로 원정대 캐릭터와 아이템 레벨을 자동으로 가져와요.",
    "   · 요일 칸 = 가능한 시간: 21-02 (21시~익일 2시), 20:30~24, 21~ (끝 생략 시 익일 02:00까지). 빈 칸 = 불가.",
    "     요일 칸을 전부 비우면 그 사람의 기존 가능 시간은 그대로 둬요.",
    "",
    "[레이드 숙제] 공대원별 캐릭터와 숙제 구간. 공대원 이름은 묶음의 첫 줄에만 적어도 돼요.",
    f"   · 숙제 구간은 드롭다운에서 선택 ('{AUTO}' = 아이템 레벨에 맞는 기본 구간).",
    "   · 직업/아이템레벨을 비우면 API로 조회해요. 오른쪽 회색 칸(레이드·클골·더보기)은 자동 계산이에요.",
    f"   · 아래로 {MAX_ROWS}줄까지 계산 칸이 미리 깔려 있어 공대원이 많아도 그대로 이어서 적으면 돼요.",
    "   · 더보기 비용은 귀속 골드가 있는 레이드는 귀속에서, 없는 레이드는 거래 가능 골드에서 빠져요.",
    "   · 실제 골드는 원정대(공대원)별 아이템 레벨 상위 6캐릭만 받아요.",
    "",
    "[일정] 한 줄 = 언제 몇 시에 어떤 캐릭터들이 같이 가는지. 같은 파티 번호 = 같이 도는 고정 파티.",
    "   · 캐릭터 칸에 같이 가는 캐릭터 이름을 쉼표로 적어요. 예: 로워붕, 캐릭터2, 캐릭터3, 캐릭터4",
    "   · 요일·시각을 적으면 매주 그 시간으로 일정이 자동으로 만들어져요.",
    "     요일만 적고 시각을 비우면 그 요일 기본 시각(config.toml 의 default_start)으로 만들어요. 나중에 /일정수정 으로 바꾸세요.",
    "     요일·시각을 둘 다 비우면 /일정제안 에서 파티원 모두 되는 시간에 파티 레이드를 이어서 제안해요.",
    "   · 파티 번호마다 줄 색이 자동으로 바뀌어요. 이 시트를 올리면 봇의 고정 파티가 시트 내용으로 통째로 바뀌어요.",
    "",
    f"회색 글씨의 예시 줄(공대원 {EX_MEMBER} / 캐릭터 {', '.join(sorted(EX_CHAR_NAMES))})은 지우고 적어 주세요. 남아 있어도 무시돼요.",
    "다 채웠으면 디스코드에서 /시트가져오기 로 이 파일을 올려 주세요.",
    "올린 뒤 미리보기를 확인하고 [적용]을 눌러야 반영돼요.",
]

HEADER_FILL = PatternFill("solid", fgColor="DDE7F7")
CALC_FILL = PatternFill("solid", fgColor="F3F3F3")
EXAMPLE_FONT = Font(color="999999", italic=True)
THIN = Side(style="thin", color="BBBBBB")

# 계산용(숨김) 시트 위치
DATA = f"'{SHEET_DATA}'"
TIER_TABLE = f"{DATA}!$A:$F"              # 구간 | 최소레벨 | 레이드1 | 레이드2 | 레이드3 | 골드X
AUTO_LEVELS, AUTO_LABELS = f"{DATA}!$I$2:$I$40", f"{DATA}!$J$2:$J$40"   # 자동 구간 (레벨 오름차순)
TIER_CHOICES_COL = "L"                    # 숙제 구간 드롭다운 (자동 + 구간들)
RAID_TABLE = f"{DATA}!$N:$S"              # 레이드 | 입장 | 인원 | 골드 | 귀속 | 더보기


# ---- 양식 만들기 --------------------------------------------------------------

def _avail_cell(slots, weekday: int) -> str:
    for a in slots:
        if a.weekday == weekday:
            s, e = fmt_range(a.start_min, a.end_min).replace("익일 ", "").split("~")
            return f"{s}-{e}"
    return ""


def build_template(db: Database, cfg: Config, guild_id: int, prefill: bool = False) -> bytes:
    """빈 양식(예시 데이터만). prefill 이면 현재 등록 정보(공대원, 가능 시간, 캐릭터·구간, 고정 파티)를 채운다."""
    wb = Workbook()
    guide = wb.active
    guide.title = SHEET_GUIDE
    for i, line in enumerate(GUIDE, 1):
        guide.cell(i, 1, line)
    guide["A1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 115

    members = db.list_members(guild_id) if prefill else []
    today = datetime.now(KST).date()
    raids_listed = sorted(
        (r for r in cfg.raids.values() if r.is_active(today)), key=lambda r: (not r.event, -r.min_item_level)
    )

    _members_sheet(wb, db, members)
    _homework_sheet(wb, db, cfg, members)
    _schedule_sheet(wb, db, cfg, guild_id if prefill else None, len(raids_listed))
    _raid_names_sheet(wb, raids_listed)
    _data_sheet(wb, cfg, raids_listed)

    wb.active = 2  # 레이드 숙제부터 열리게
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _members_sheet(wb, db: Database, members) -> None:
    ws = wb.create_sheet(SHEET_MEMBERS)
    _header(ws, MEMBER_HEADERS, [14, 16] + [12] * 7)
    if members:
        for m in members:
            slots = db.list_availability(m.id)
            ws.append([m.display_name, m.roster_name or ""] + [_avail_cell(slots, d) for d in range(7)])
    else:
        _example(ws, [EX_MEMBER, EX_CHARS[0][0], "", "", "21-02", "20-24", "21~", "14-20", ""])
    for c in range(1, len(MEMBER_HEADERS) + 1):
        _text_column(ws, c)  # 시간이 날짜/시각 값으로 바뀌지 않게 텍스트로


def _homework_sheet(wb, db: Database, cfg: Config, members) -> None:
    ws = wb.create_sheet(SHEET_HOMEWORK)
    headers = HOMEWORK_HEADERS + HOMEWORK_CALC
    _header(ws, headers, [12, 16, 11, 11, 11, 11, 17, 17, 17, 17, 11, 11, 11])
    for c in range(6, len(headers) + 1):
        ws.cell(1, c).fill = PatternFill("solid", fgColor="E7E6E6")

    row = 2
    for m in members:
        chars = db.list_characters(m.id)
        for j, c in enumerate(chars):
            tier = c.tier if c.tier in cfg.tiers else AUTO
            ws.cell(row, 1, m.display_name if j == 0 else None)
            for col, v in enumerate([c.name, c.cls, f"{c.item_level:g}", tier], 2):
                ws.cell(row, col, v)
            row += 1
        if chars:
            ws.cell(row - len(chars), 1).font = Font(bold=True)
            for col in range(1, len(headers) + 1):  # 공대원 묶음 사이 구분선
                ws.cell(row - 1, col).border = Border(bottom=THIN)
    if row == 2:
        for j, (name, cls, level, tier) in enumerate(EX_CHARS):
            for col, v in enumerate([EX_MEMBER if j == 0 else None, name, cls, level, tier], 1):
                if v is not None:
                    ws.cell(row, col, v).font = EXAMPLE_FONT
            row += 1

    # 계산 칸은 MAX_ROWS 줄까지 미리 깔아 둔다 (사람이 아래로 계속 적어도 계산되게)
    for r in range(2, MAX_ROWS + 1):
        _homework_formulas(ws, r)
    _text_column(ws, 4)
    dv = DataValidation(
        type="list", formula1=f"{DATA}!${TIER_CHOICES_COL}$2:${TIER_CHOICES_COL}${len(cfg.tiers) + 2}",
        allow_blank=True,
    )
    ws.add_data_validation(dv)
    dv.add(f"E2:E{MAX_ROWS}")


def _homework_formulas(ws, r: int) -> None:
    # F 적용 구간: '자동'(또는 빈칸)이면 아이템 레벨로 자동 구간
    ws.cell(r, 6, (
        f'=IF($B{r}="","",IF(OR($E{r}="{AUTO}",$E{r}=""),'
        f'IF($D{r}="","",IFERROR(LOOKUP(VALUE($D{r}),{AUTO_LEVELS},{AUTO_LABELS}),"")),$E{r}))'
    ))
    for i, col in enumerate("GHIJ"):  # 레이드1~3, 골드X ← 구간표 C~F. &"" : 빈 칸이 0 으로 나오지 않게
        ws[f"{col}{r}"] = f'=IF($F{r}="","",IFERROR(VLOOKUP($F{r},{TIER_TABLE},{i + 3},FALSE)&"",""))'
    gold = "+".join(f'IFERROR(VLOOKUP({c}{r},{RAID_TABLE},4,FALSE),0)' for c in "GHI")
    more = "+".join(f'IFERROR(VLOOKUP({c}{r},{RAID_TABLE},6,FALSE),0)' for c in "GHI")
    ws[f"K{r}"] = f'=IF($F{r}="","",{gold})'
    ws[f"L{r}"] = f'=IF($F{r}="","",{more})'
    ws[f"M{r}"] = f'=IF($F{r}="","",K{r}-L{r})'
    for col in range(6, 14):
        ws.cell(r, col).fill = CALC_FILL
    for col in "KLM":
        ws[f"{col}{r}"].number_format = "#,##0"


def _schedule_sheet(wb, db: Database, cfg: Config, guild_id: int | None, raid_count: int) -> None:
    """고정 파티 일정: 한 줄 = 언제 몇 시에 어떤 캐릭터들이 같이 가는지.
    파티 번호별로 줄 색이 바뀐다 (조건부 서식이라 새로 추가한 파티에도 적용)."""
    ws = wb.create_sheet(SHEET_SCHEDULE)
    headers = PARTY_FIXED + [PARTY_CHARS]
    _header(ws, headers, [6, 9, 20, 6, 8, 70])

    parties = db.list_parties(guild_id) if guild_id is not None else []
    for p in parties:
        for r in p.raids:
            raid = cfg.raid(r.raid_key)
            row = [p.number, p.label or "", raid.name if raid else r.raid_key]
            row += [WEEKDAYS[r.weekday], fmt_minutes(r.start_min).replace("익일 ", "")] if r.timed else ["", ""]
            ws.append(row + [", ".join(c.name for c in r.characters)])
    if not parties:
        for values in EX_SCHEDULE:
            _example(ws, list(values))

    last_col = ws.cell(1, len(headers)).column_letter
    area = f"A2:{last_col}{MAX_ROWS}"
    for k, color in enumerate(PARTY_FILLS):
        ws.conditional_formatting.add(area, FormulaRule(
            formula=[f"AND(ISNUMBER($A2),MOD($A2-1,{len(PARTY_FILLS)})={k})"],
            fill=PatternFill("solid", start_color=color, end_color=color),
        ))
    _text_column(ws, 5)
    raid_dv = DataValidation(
        type="list", formula1=f"='{SHEET_RAID_NAMES}'!$A$2:$A${raid_count + 1}", allow_blank=True
    )
    day_dv = DataValidation(type="list", formula1='"' + ",".join(WEEKDAYS) + '"', allow_blank=True)
    tier_dv = DataValidation(type="list", formula1=f"{DATA}!$A$2:$A$40", allow_blank=True)
    for dv, rng in ((raid_dv, f"C2:C{MAX_ROWS}"), (day_dv, f"D2:D{MAX_ROWS}"), (tier_dv, f"B2:B{MAX_ROWS}")):
        ws.add_data_validation(dv)
        dv.add(rng)


def _raid_names_sheet(wb, raids) -> None:
    ws = wb.create_sheet(SHEET_RAID_NAMES)
    _header(ws, ["레이드", "입장 레벨", "인원"], [24, 10, 6])
    for r in raids:
        ws.append([r.name, r.min_item_level, r.party_size])


def _data_sheet(wb, cfg: Config, raids) -> None:
    """레이드 숙제 수식이 참조하는 표. 사람이 볼 필요 없어 숨긴다 (구간표는 config.toml 에서 관리)."""
    ws = wb.create_sheet(SHEET_DATA)
    ws.sheet_state = "hidden"
    ws.append(["구간", "최소 레벨", "레이드1", "레이드2", "레이드3", "골드X"])
    tiers = sorted(cfg.tiers.values(), key=lambda t: (-t.min_level, t.label))
    for i, t in enumerate(tiers, 2):
        names = [cfg.raids[k].name for k in t.raids] + [""] * (3 - len(t.raids))
        for col, v in enumerate([t.label, t.min_level, *names[:3], cfg.raids[t.no_gold].name if t.no_gold else ""], 1):
            ws.cell(i, col, v)
    ws["I1"], ws["J1"], ws[f"{TIER_CHOICES_COL}1"] = "자동 최소 레벨", "자동 구간", "선택지"
    for i, t in enumerate(sorted((t for t in tiers if t.auto), key=lambda t: t.min_level), 2):
        ws.cell(i, 9, t.min_level)
        ws.cell(i, 10, t.label)
    ws[f"{TIER_CHOICES_COL}2"] = AUTO
    for i, t in enumerate(tiers, 3):
        ws[f"{TIER_CHOICES_COL}{i}"] = t.label
    for col, v in enumerate(["레이드", "입장 레벨", "인원", "골드", "귀속", "더보기"], 14):
        ws.cell(1, col, v)
    for i, r in enumerate(raids, 2):
        for col, v in enumerate([r.name, r.min_item_level, r.party_size, r.gold, r.bound_gold, r.more_cost], 14):
            ws.cell(i, col, v)


def _header(ws, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        cell = ws.cell(1, i)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"


def _text_column(ws, col: int, last: int = MAX_ROWS) -> None:
    """입력값이 날짜/숫자로 자동 변환되지 않게 텍스트 서식. 데이터를 다 쓴 뒤 호출할 것
    (서식만 있는 빈 칸도 '사용 중'으로 잡혀 이후 append 가 그 아래로 밀리기 때문)."""
    for row in range(2, last + 1):
        ws.cell(row, col).number_format = "@"


def _example(ws, values: list) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font = EXAMPLE_FONT


# ---- 읽기 ---------------------------------------------------------------------

@dataclass
class MemberInput:
    row: int
    name: str
    discord_id: int | None
    roster_name: str | None
    # None 이면 가능 시간을 건드리지 않음. dict 면 요일별 (시작, 끝), 빠진 요일 = 불가
    availability: dict[int, tuple[int, int]] | None


@dataclass
class CharacterInput:
    row: int
    member_name: str
    name: str
    cls: str | None
    item_level: float | None
    tier: str | None = None   # None = 자동
    tier_given: bool = False  # 숙제 구간 칸이 있는 시트에서 읽었는지 (이전 양식이면 구간을 건드리지 않음)


@dataclass
class SessionInput:
    row: int
    start: datetime
    raid_key: str
    character_names: list[str]
    memo: str | None


@dataclass
class PartyRowInput:
    row: int
    raid_key: str
    weekday: int | None
    start_min: int | None
    entries: list[tuple[str | None, str]]  # (공대원 열 머리글 — 이전 양식만, 캐릭터 이름)


@dataclass
class PartyInput:
    number: int
    label: str | None
    rows: list[PartyRowInput] = field(default_factory=list)


@dataclass
class SheetData:
    members: list[MemberInput] = field(default_factory=list)
    characters: list[CharacterInput] = field(default_factory=list)
    sessions: list[SessionInput] = field(default_factory=list)
    # None = 파티 시트 없음(기존 파티 유지). [] = 파티 시트가 비어 있음(파티 전부 삭제)
    parties: list[PartyInput] | None = None
    party_members: list[str] = field(default_factory=list)  # 이전 양식 파티 시트의 공대원 열 머리글
    errors: list[str] = field(default_factory=list)

    def member_names(self) -> list[str]:
        """시트 어디서든 언급된 공대원 이름 (중복 제거, 등장 순)."""
        used_party = {m for p in self.parties or [] for r in p.rows for m, _ in r.entries}
        names = [m.name for m in self.members] + [c.member_name for c in self.characters]
        names += [m for m in self.party_members if m in used_party]
        return list(dict.fromkeys(names))


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _header_index(ws) -> dict[str, int]:
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    out: dict[str, int] = {}
    for i, v in enumerate(row):
        name = _text(v)
        if name and name not in out:
            out[name] = i
    return out


def _rows(ws, headers: list[str]):
    """(행 번호, {헤더: 값}) — 빈 줄과 예시 줄은 건너뛴다. 열 순서가 바뀌어도 헤더 이름으로 찾는다."""
    index = {h: i for h, i in _header_index(ws).items() if h in headers}
    # 서식만 지정된 빈 줄이 있을 수 있어 끝까지 읽고 빈 줄은 건너뛴다
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = {h: (row[i] if i < len(row) else None) for h, i in index.items()}
        if all(_text(v) == "" for v in values.values()):
            continue
        if any(_text(v).startswith(EXAMPLE_PREFIX) for v in values.values()):
            continue
        yield r, values


def _missing_headers(ws, headers: list[str]) -> list[str]:
    present = _header_index(ws)
    return [h for h in headers if h not in present]


def parse_workbook(content: bytes, cfg: Config, today: date) -> SheetData:
    data = SheetData()
    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as e:  # 손상된 파일, xlsx 가 아닌 파일 등
        data.errors.append(f"엑셀 파일을 열 수 없어요 ({type(e).__name__}). /시트양식 으로 받은 .xlsx 파일을 올려 주세요.")
        return data

    sheets = set(wb.sheetnames)
    known = {SHEET_MEMBERS, SHEET_HOMEWORK, SHEET_CHARS_OLD, SHEET_SCHEDULE, SHEET_PARTY_OLD}
    if not sheets & known:
        data.errors.append("공대원/레이드 숙제/일정 시트를 찾지 못했어요. /시트양식 으로 받은 양식을 써 주세요.")
        return data

    if SHEET_MEMBERS in sheets:
        _parse_members(wb[SHEET_MEMBERS], data)
    if SHEET_HOMEWORK in sheets:
        _parse_homework(wb[SHEET_HOMEWORK], cfg, data)
    elif SHEET_CHARS_OLD in sheets:
        _parse_homework(wb[SHEET_CHARS_OLD], cfg, data, old=True)
    schedule_is_party = SHEET_SCHEDULE in sheets and "파티" in _header_index(wb[SHEET_SCHEDULE])
    if schedule_is_party:
        _parse_parties(wb[SHEET_SCHEDULE], cfg, data)
    elif SHEET_PARTY_OLD in sheets:
        _parse_parties(wb[SHEET_PARTY_OLD], cfg, data)
    if SHEET_SCHEDULE in sheets and not schedule_is_party:
        _parse_sessions(wb[SHEET_SCHEDULE], cfg, today, data)
    wb.close()
    return data


def _parse_members(ws, data: SheetData) -> None:
    if missing := _missing_headers(ws, ["공대원"]):
        data.errors.append(f"[공대원] 머리글 '{missing[0]}' 이 없어요.")
        return
    seen: set[str] = set()
    for r, v in _rows(ws, MEMBER_HEADERS + ["디스코드ID"]):
        where = f"[공대원 {r}행]"
        name = _text(v.get("공대원"))
        if not name:
            data.errors.append(f"{where} 공대원 이름이 비어 있어요.")
            continue
        if name in seen:
            data.errors.append(f"{where} '{name}' 이 두 번 나와요.")
            continue
        seen.add(name)

        raw_id = _text(v.get("디스코드ID"))
        discord_id = None
        if raw_id:
            if not re.fullmatch(r"\d{17,20}", raw_id):
                data.errors.append(f"{where} 디스코드ID '{raw_id}' 가 올바르지 않아요 (숫자 17~20자리, 텍스트로 입력).")
                continue
            discord_id = int(raw_id)

        avail: dict[int, tuple[int, int]] = {}
        filled = False
        bad = False
        for d, day in enumerate(WEEKDAYS):
            text = _time_cell_text(v.get(day))
            if not text:
                continue
            filled = True
            if text in ("x", "X", "불가", "-"):
                continue
            try:
                avail[d] = parse_time_range(text)
            except ParseError:
                data.errors.append(f"{where} {day}요일 '{text}' 를 이해하지 못했어요. 예: 21-02, 20:30~24, 21~")
                bad = True
        if bad:
            continue
        roster = _text(v.get("대표캐릭터")) or None
        if name == EX_MEMBER and roster in EX_CHAR_NAMES:
            continue  # 예시 줄
        data.members.append(MemberInput(r, name, discord_id, roster, avail if filled else None))


def _time_cell_text(cell) -> str:
    # 엑셀이 '21:00' 을 시각 값으로 바꿔 버린 경우
    if isinstance(cell, (time, datetime)):
        return cell.strftime("%H:%M")
    return _text(cell)


def _parse_homework(ws, cfg: Config, data: SheetData, old: bool = False) -> None:
    sheet = SHEET_CHARS_OLD if old else SHEET_HOMEWORK
    name_header = "캐릭터명" if old else "캐릭터"
    if missing := _missing_headers(ws, ["공대원", name_header]):
        data.errors.append(f"[{sheet}] 머리글 '{missing[0]}' 이 없어요.")
        return
    has_tier = "숙제 구간" in _header_index(ws)
    member = ""
    for r, v in _rows(ws, ["공대원", name_header, "직업", "아이템레벨", "숙제 구간"]):
        where = f"[{sheet} {r}행]"
        member = _text(v.get("공대원")) or member  # 공대원은 묶음 첫 줄에만 적어도 된다
        name = _text(v.get(name_header))
        if not name or name == TOTAL_ROW or (member == EX_MEMBER and name in EX_CHAR_NAMES):
            continue  # 빈 줄, 이전 양식 합계 줄, 예시 줄
        if not member:
            data.errors.append(f"{where} {name} 의 공대원이 없어요.")
            continue
        level_raw = _text(v.get("아이템레벨")).replace(",", "")
        level = None
        if level_raw:
            try:
                level = float(level_raw)
            except ValueError:
                data.errors.append(f"{where} 아이템레벨 '{level_raw}' 가 숫자가 아니에요.")
                continue
        tier_text = _text(v.get("숙제 구간"))
        tier = None if tier_text in ("", AUTO) else tier_text
        if tier and tier not in cfg.tiers:
            data.errors.append(f"{where} 숙제 구간 '{tier_text}' 이 구간표에 없어요. 자동으로 둘게요.")
            tier = None
        data.characters.append(
            CharacterInput(r, member, name, _text(v.get("직업")) or None, level, tier, tier_given=has_tier)
        )


def find_raid_key(cfg: Config, text: str) -> str | None:
    norm = text.replace(" ", "").removeprefix(EVENT_PREFIX.strip())
    for r in cfg.raids.values():
        if norm in (r.name.replace(" ", "").removeprefix(EVENT_PREFIX.strip()), r.key):
            return r.key
    return None


def _parse_parties(ws, cfg: Config, data: SheetData) -> None:
    if missing := _missing_headers(ws, ["파티", "레이드"]):
        data.errors.append(f"[{ws.title}] 머리글 '{missing[0]}' 이 없어요. 파티 일정은 건너뛸게요.")
        return
    header = _header_index(ws)
    fixed = set(PARTY_FIXED) | LEGACY_NOTE | {PARTY_CHARS}
    # 이전 양식: 공대원 이름 열마다 캐릭터 하나. 지금 양식: '캐릭터' 칸 하나에 여러 이름
    member_cols = [(h, i) for h, i in sorted(header.items(), key=lambda kv: kv[1]) if h not in fixed]
    data.party_members = [h for h, _ in member_cols]
    parties: dict[int, PartyInput] = {}

    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        get = lambda h: row[header[h]] if h in header and header[h] < len(row) else None  # noqa: E731
        number_text = _text(get("파티"))
        raid_text = _text(get("레이드"))
        if not number_text and not raid_text:
            continue
        if number_text.startswith(EXAMPLE_PREFIX):
            continue
        where = f"[{ws.title} {r}행]"
        if not number_text.isdigit() or int(number_text) < 1:
            data.errors.append(f"{where} 파티 번호 '{number_text}' 는 1 이상의 숫자여야 해요.")
            continue
        raid_key = find_raid_key(cfg, raid_text) if raid_text not in SKIP else None
        if raid_key is None:
            if raid_text not in SKIP:
                data.errors.append(f"{where} 레이드 '{raid_text}' 를 찾을 수 없어요. 드롭다운에서 골라 주세요.")
            continue

        weekday = start_min = None
        day_text, time_text = _text(get("요일")).removesuffix("요일"), _time_cell_text(get("시각"))
        if day_text or time_text:
            if day_text not in WEEKDAYS:
                data.errors.append(f"{where} 시각만 있고 요일이 없어요. 요일도 적어 주세요 (예: 수 / 21:00). 시간 미정으로 둘게요.")
            else:
                try:
                    # 요일만 적고 시각을 비우면 기본 시각으로 그 요일에 파티를 만든다
                    t = parse_time(time_text or cfg.default_start)
                    weekday, start_min = WEEKDAYS.index(day_text), t.hour * 60 + t.minute
                except ParseError as e:
                    data.errors.append(f"{where} {e} 시간 미정으로 둘게요.")

        entries: list[tuple[str | None, str]] = []
        for name in re.split(r"[,\n/]", _text(get(PARTY_CHARS))):
            if name.strip() not in SKIP:
                entries.append((None, name.strip()))
        for member, i in member_cols:
            name = _text(row[i]) if i < len(row) else ""
            if name not in SKIP:
                entries.append((member, name))
        if entries and all(n in EX_CHAR_NAMES for _, n in entries):
            continue  # 예시 줄
        number = int(number_text)
        party = parties.setdefault(number, PartyInput(number, _text(get("구간")) or None))
        if party.label is None and _text(get("구간")):
            party.label = _text(get("구간"))
        party.rows.append(PartyRowInput(r, raid_key, weekday, start_min, entries))
    data.parties = [parties[n] for n in sorted(parties)]


def _parse_sessions(ws, cfg: Config, today: date, data: SheetData) -> None:
    if missing := _missing_headers(ws, ["날짜", "시각", "레이드"]):
        data.errors.append(f"[일정] 머리글 '{missing[0]}' 이 없어요.")
        return
    for r, v in _rows(ws, SESSION_HEADERS):
        where = f"[일정 {r}행]"
        try:
            day_cell, time_cell = v.get("날짜"), v.get("시각")
            if isinstance(day_cell, datetime):
                day = day_cell.date()
            elif isinstance(day_cell, date):
                day = day_cell
            else:
                day = parse_date(_text(day_cell), today)
            at = time_cell if isinstance(time_cell, time) else parse_time(_time_cell_text(time_cell))
        except ParseError as e:
            data.errors.append(f"{where} {e}")
            continue
        raid_text = _text(v.get("레이드"))
        raid_key = find_raid_key(cfg, raid_text)
        if raid_key is None:
            data.errors.append(f"{where} 레이드 '{raid_text}' 를 찾을 수 없어요. 드롭다운에서 골라 주세요.")
            continue
        names = [n.strip() for n in re.split(r"[,\n/]", _text(v.get("참가 캐릭터"))) if n.strip()]
        data.sessions.append(
            SessionInput(r, datetime.combine(day, at.replace(tzinfo=None), KST), raid_key, names,
                         _text(v.get("메모")) or None)
        )
